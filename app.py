import streamlit as st
import pandas as pd
import os
import shutil
import json
from PyPDF2 import PdfReader
from docx import Document
from fpdf import FPDF

# Configuration
BASE_FOLDER = "data"
RECRUITED_FOLDER = os.path.join(BASE_FOLDER, "retenus")
OUTPUT_FILE = os.path.join(BASE_FOLDER, "classement_candidats.csv")
KEYWORDS_FILE = os.path.join(BASE_FOLDER, "keywords.json")
MIN_KEYWORDS = 5

# CSS pour interface mobile-friendly
st.markdown("""
    <style>
        .block-container {
            padding-top: 1rem;
            padding-bottom: 2rem;
        }
        .stTextInput>div>div>input,
        .stTextArea>div>textarea,
        .stSelectbox>div>div>div {
            font-size: 1rem;
        }
        .stButton>button {
            font-size: 1rem;
            padding: 0.5rem 1.25rem;
        }
    </style>
""", unsafe_allow_html=True)

# Extraction de texte
@st.cache_data
def extract_text(path):
    try:
        if path.endswith(".pdf"):
            return " ".join([p.extract_text() or "" for p in PdfReader(path).pages]).lower()
        elif path.endswith(".docx"):
            return " ".join(p.text for p in Document(path).paragraphs).lower()
    except:
        return ""

# Comptage de mots-clés
def count_keywords(text, keywords):
    return sum(1 for kw in keywords if kw in text)

# Génération de résumé PDF
def generate_pdf_summary(result):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, txt="Résumé de candidature", ln=True, align="C")
    pdf.ln(10)

    for key in ["Nom", "Profil", "Score"]:
        pdf.cell(200, 10, txt=f"{key} : {result[key]}", ln=True)

    pdf.ln(5)
    pdf.multi_cell(0, 10, txt=f"Extrait du CV :\n{result['Extrait']}")

    summary_folder = os.path.join(RECRUITED_FOLDER, result['Profil'], "web_upload")
    os.makedirs(summary_folder, exist_ok=True)
    pdf_path = os.path.join(summary_folder, f"resume_{result['Nom'].replace(' ', '_')}.pdf")
    pdf.output(pdf_path)
    return pdf_path

# Traitement des fichiers téléversés
def process_uploaded_files(uploaded_files, profile, keywords):
    results = []
    temp_folder = os.path.join(BASE_FOLDER, profile, "web_upload")
    os.makedirs(temp_folder, exist_ok=True)

    for uploaded_file in uploaded_files:
        filepath = os.path.join(temp_folder, uploaded_file.name)
        with open(filepath, "wb") as f:
            f.write(uploaded_file.getbuffer())

        text = extract_text(filepath)
        score = count_keywords(text, keywords)

        result = {
            "Nom": uploaded_file.name,
            "Profil": profile,
            "Score": score,
            "Extrait": text[:500],
            "Texte complet": text,
            "Fichier": filepath
        }
        results.append(result)

        if score >= MIN_KEYWORDS:
            dest_folder = os.path.join(RECRUITED_FOLDER, profile, "web_upload")
            os.makedirs(dest_folder, exist_ok=True)
            shutil.copy(filepath, os.path.join(dest_folder, uploaded_file.name))
            generate_pdf_summary(result)

    return results

# Chargement / sauvegarde des mots-clés enregistrés
def load_saved_keywords():
    if os.path.exists(KEYWORDS_FILE):
        with open(KEYWORDS_FILE, "r") as f:
            return json.load(f)
    return {}

def save_keywords(profiles_keywords):
    with open(KEYWORDS_FILE, "w") as f:
        json.dump(profiles_keywords, f, indent=2)

# Interface Streamlit
st.set_page_config(page_title="Analyse de CV", layout="centered")
st.title("📋 Analyse de CV multi-profils")
st.markdown("\nTéléversez les CVs, gérez vos mots-clés par profil, et effectuez des recherches plein texte.")

saved_keywords = load_saved_keywords()

with st.expander("🔍 Editeur de mots-clés par profil"):
    selected_profile = st.text_input("Nom du profil pour édition", value=list(saved_keywords.keys())[0] if saved_keywords else "")
    existing_keywords = ", ".join(saved_keywords.get(selected_profile, []))
    new_keywords_input = st.text_area("Mots-clés (séparés par des virgules)", existing_keywords)
    if st.button("Enregistrer les mots-clés"):
        saved_keywords[selected_profile] = [kw.strip().lower() for kw in new_keywords_input.split(",") if kw.strip()]
        save_keywords(saved_keywords)
        st.success("Mots-clés enregistrés.")

# Choix du profil pour analyse
profile = st.selectbox("Choisissez un profil pour analyse", list(saved_keywords.keys()))
uploaded_files = st.file_uploader("Téléverser des CV (.pdf ou .docx)", type=["pdf", "docx"], accept_multiple_files=True)

if profile and uploaded_files:
    keywords = saved_keywords.get(profile, [])
    with st.spinner("Analyse en cours..."):
        results = process_uploaded_files(uploaded_files, profile, keywords)
        df = pd.DataFrame(results)
        df.to_csv(OUTPUT_FILE, index=False)
        st.success(f"{len(results)} fichier(s) analysé(s). Rapport enregistré dans {OUTPUT_FILE}")
        st.dataframe(df[["Nom", "Profil", "Score"]])

        if not df.empty:
            st.download_button("📥 Télécharger le classement CSV", df.to_csv(index=False), file_name="classement_candidats.csv", mime="text/csv")

            with st.expander("📂 Aperçu texte brut (500 premiers caractères)"):
                for i, row in df.iterrows():
                    st.markdown(f"**{row['Nom']}** - Score: {row['Score']}")
                    st.code(row['Extrait'], language='text')

            with st.expander("🔎 Recherche plein texte dans tous les CV"):
                query = st.text_input("Rechercher un mot ou une phrase :")
                if query:
                    query = query.lower()
                    matches = df[df['Texte complet'].str.contains(query, na=False)]
                    if not matches.empty:
                        for _, row in matches.iterrows():
                            st.markdown(f"**{row['Nom']}** (Score: {row['Score']})")
                            st.code(row['Texte complet'][:1000])
                    else:
                        st.warning("Aucun CV ne contient ce terme.")
